#!/usr/bin/env python3
"""
postmortem_docx.py — styled .docx builder for engagement post-mortem / read-out docs.

House style (matches the SA's Partner Oversight Read-Out / Post-Mortem template):
  - Calibri throughout
  - Navy #1F3864 title + section headers; gray #6B7280 subtitle/captions; body #444B54
  - Key-value tables: shaded #F2F4F7 label cells
  - Data tables: navy header row, white text
  - US Letter, 0.8in top/bottom + 0.9in left/right margins, 1.12 line spacing, no header/footer

Usage:
    from postmortem_docx import PostMortemDoc
    d = PostMortemDoc()
    d.title("Acme — Partner Oversight Read-Out")
    d.subtitle("Retrospective Risk Assessment · UAT Answer · Design-Review Gap")
    d.meta("Interim · current as of June 2, 2026 · delivered through Partner · prepared for Services leadership")
    d.h1("Where Things Stand"); d.caption("As of June 2, 2026.")
    d.kv_table([("Status", "ESCALATED (red) — ..."), ("Top risks", "...")])
    d.h1("Executive Summary"); d.body("...")
    d.data_table(["#", "Event", "Impact"], [["1", "...", "..."]], widths=(0.35, 4.05, 3.0))
    d.body([("Bold lead — ", True), ("rest of sentence.", False)])   # mixed runs
    d.save("/path/Acme_Read-Out.docx")

Requires: python-docx  (pip install python-docx --break-system-packages)
"""
from docx import Document
from docx.shared import Pt, RGBColor, Twips
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
GRAY = RGBColor(0x6B, 0x72, 0x80)
BODY = RGBColor(0x44, 0x4B, 0x54)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LABEL_FILL = "F2F4F7"
HEADER_FILL = "1F3864"
BORDER = "D0D5DD"
FONT = "Calibri"


class PostMortemDoc:
    def __init__(self):
        self.doc = Document()
        normal = self.doc.styles["Normal"]
        normal.font.name = FONT
        normal.font.size = Pt(11)
        normal.font.color.rgb = BODY
        sec = self.doc.sections[0]
        sec.page_width = Twips(12240)
        sec.page_height = Twips(15840)
        sec.top_margin = Twips(1152)
        sec.bottom_margin = Twips(1152)
        sec.left_margin = Twips(1296)
        sec.right_margin = Twips(1296)
        sec.header_distance = Twips(720)
        sec.footer_distance = Twips(720)

    # ---- paragraph primitives ----
    def _para(self, runs, size=11, color=BODY, bold=False, before=0, after=6, italic=False):
        p = self.doc.add_paragraph()
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_before = Pt(before)
        p.paragraph_format.space_after = Pt(after)
        if isinstance(runs, str):
            runs = [(runs, bold)]
        for text, b in runs:
            r = p.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = color
            r.font.bold = b
            r.font.italic = italic
        return p

    def title(self, text):
        """Document title — navy, 19pt bold."""
        return self._para(text, size=19, color=NAVY, bold=True, after=2)

    def subtitle(self, text):
        """Pillar line under the title — gray, 10.5pt. Use '·' separators."""
        return self._para(text, size=10.5, color=GRAY, after=1)

    def meta(self, text):
        """Status/provenance line — gray, 9.5pt. e.g. 'Interim · current as of ... · prepared for ...'."""
        return self._para(text, size=9.5, color=GRAY, after=6)

    def h1(self, text):
        """Section header — navy, 14pt bold."""
        return self._para(text, size=14, color=NAVY, bold=True, before=12, after=4)

    def h2(self, text):
        """Sub-header — navy, 11.5pt bold."""
        return self._para(text, size=11.5, color=NAVY, bold=True, before=8, after=2)

    def caption(self, text):
        """Gray sub-caption under a section header — 10pt."""
        return self._para(text, size=10, color=GRAY, after=4)

    def body(self, runs, before=0, after=6):
        """Body paragraph. Pass a string, or a list of (text, bold) tuples for inline bolding."""
        return self._para(runs, size=11, color=BODY, before=before, after=after)

    def spacer(self):
        p = self.doc.add_paragraph()
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(2)
        return p

    # ---- table helpers ----
    def _shade(self, cell, fill):
        tcPr = cell._tc.get_or_add_tcPr()
        shd = OxmlElement("w:shd")
        shd.set(qn("w:val"), "clear")
        shd.set(qn("w:fill"), fill)
        tcPr.append(shd)

    def _borders(self, table):
        tblPr = table._tbl.tblPr
        borders = OxmlElement("w:tblBorders")
        for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
            e = OxmlElement(f"w:{edge}")
            e.set(qn("w:val"), "single")
            e.set(qn("w:sz"), "4")
            e.set(qn("w:space"), "0")
            e.set(qn("w:color"), BORDER)
            borders.append(e)
        tblPr.append(borders)

    def _cell(self, cell, runs, bold=False, color=BODY, size=10, white=False):
        cell.text = ""
        p = cell.paragraphs[0]
        p.paragraph_format.line_spacing = 1.12
        p.paragraph_format.space_before = Pt(2)
        p.paragraph_format.space_after = Pt(2)
        if isinstance(runs, str):
            runs = [(runs, bold)]
        for text, b in runs:
            r = p.add_run(text)
            r.font.name = FONT
            r.font.size = Pt(size)
            r.font.color.rgb = WHITE if white else color
            r.font.bold = b

    def kv_table(self, rows, widths=(2.0, 5.4)):
        """Two-column key-value table (status block). rows = [(label, value), ...]."""
        t = self.doc.add_table(rows=0, cols=2)
        self._borders(t)
        for label, value in rows:
            tr = t.add_row().cells
            self._cell(tr[0], label, bold=True, color=NAVY, size=10)
            self._shade(tr[0], LABEL_FILL)
            self._cell(tr[1], value, color=BODY, size=10)
        for r in t.rows:
            r.cells[0].width = Pt(widths[0] * 72)
            r.cells[1].width = Pt(widths[1] * 72)
        return t

    def data_table(self, headers, rows, widths):
        """Multi-column data table with a navy header row. widths = column widths in inches."""
        t = self.doc.add_table(rows=0, cols=len(headers))
        self._borders(t)
        hr = t.add_row().cells
        for i, h in enumerate(headers):
            self._cell(hr[i], h, bold=True, white=True, size=10)
            self._shade(hr[i], HEADER_FILL)
        for row in rows:
            rc = t.add_row().cells
            for i, val in enumerate(row):
                self._cell(rc[i], val, color=BODY, size=10)
        for r in t.rows:
            for i, w in enumerate(widths):
                r.cells[i].width = Pt(w * 72)
        return t

    def save(self, path):
        self.doc.save(path)
        return path
